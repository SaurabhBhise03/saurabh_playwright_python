import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as OpenpyxlImage

class ExcelReporter:
    """
    Process-isolated Excel logging manager designed to prevent file corruption
    and object-copy serialization crashes across parallel worker nodes.
    """
    def __init__(self, execution_dir: str, worker_id: str = "master"):
        self.target_path = os.path.join(execution_dir, f"Summary_{worker_id}.xlsx")
        self._initialize_workbook()

    def _initialize_workbook(self):
        """Initializes individual worker spreadsheets with custom executive headers."""
        if not os.path.exists(self.target_path):
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Detailed Step Logs"
            
            # Column E tracks the raw path string so the merge hook can locate the image file on disk safely
            headers = ["Test Case Name", "Step Description", "Status Outcome", "Duration (s)", "Screenshot Path File"]
            sheet.append(headers)
            
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            
            column_widths = {"A": 35, "B": 55, "C": 15, "D": 15, "E": 55}
            for col, width in column_widths.items():
                sheet.column_dimensions[col].width = width

            for col_idx in range(1, 6):
                cell = sheet.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
            workbook.save(self.target_path)

    def append_step_result(self, test_name: str, step_desc: str, status: str, duration: float, screenshot_path: str = None):
        """Appends validation steps directly to the isolated worker's excel workbook file."""
        workbook = load_workbook(self.target_path)
        sheet = workbook.active
        
        # Save the screenshot path as text inside column E
        path_to_log = screenshot_path if screenshot_path else "N/A"
        sheet.append([test_name, step_desc, status.upper(), round(duration, 2), path_to_log])
        current_row = sheet.max_row
        
        for col_idx in range(1, 6):
            sheet.cell(row=current_row, column=col_idx).alignment = Alignment(vertical="center", wrap_text=True)

        status_cell = sheet.cell(row=current_row, column=3)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        if status.upper() == "PASSED":
            status_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            status_font = Font(name="Segoe UI", size=10, bold=True, color="006100")
        else:
            status_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            status_font = Font(name="Segoe UI", size=10, bold=True, color="9C0006")
            
        status_cell.fill = status_fill
        status_cell.font = status_font

        workbook.save(self.target_path)

    @staticmethod
    def merge_worker_reports(execution_dir: str, output_filename: str):
        """Consolidates individual worker log segments into a single master summary report with thumbnail grids."""
        master_path = os.path.join(execution_dir, output_filename)
        master_wb = Workbook()
        master_sheet = master_wb.active
        master_sheet.title = "Consolidated Logs"
        
        headers = ["Test Case Name", "Step Description", "Status Outcome", "Duration (s)", "Step Screenshot Evidence"]
        master_sheet.append(headers)
        
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        
        column_widths = {"A": 35, "B": 55, "C": 15, "D": 15, "E": 45}
        for col, width in column_widths.items():
            master_sheet.column_dimensions[col].width = width
        for col_idx in range(1, 6):
            cell = master_sheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Read segments and construct the master workbook cleanly
        for file in os.listdir(execution_dir):
            if file.startswith("Summary_") and file.endswith(".xlsx"):
                file_path = os.path.join(execution_dir, file)
                worker_wb = load_workbook(file_path)
                worker_sheet = worker_wb.active
                
                for row_idx in range(2, worker_sheet.max_row + 1):
                    row_values = [worker_sheet.cell(row=row_idx, column=c).value for c in range(1, 5)]
                    img_path_value = worker_sheet.cell(row=row_idx, column=5).value
                    
                    master_sheet.append(row_values + [""])
                    m_row = master_sheet.max_row
                    
                    # Status styling extraction lookup logic
                    s_cell = master_sheet.cell(row=m_row, column=3)
                    s_cell.alignment = Alignment(horizontal="center", vertical="center")
                    if str(s_cell.value).upper() == "PASSED":
                        s_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                        s_cell.font = Font(name="Segoe UI", size=10, bold=True, color="006100")
                    else:
                        s_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                        s_cell.font = Font(name="Segoe UI", size=10, bold=True, color="9C0006")

                    # Draw screenshot thumbnail dynamically using the absolute text path location
                    if img_path_value and os.path.exists(img_path_value):
                        try:
                            img = OpenpyxlImage(img_path_value)
                            img.width = 320
                            img.height = 180
                            master_sheet.row_dimensions[m_row].height = 140
                            master_sheet.add_image(img, f"E{m_row}")
                        except Exception:
                            master_sheet.cell(row=m_row, column=5).value = "Screenshot Available (Error Rendering Thumbnail)"
                    else:
                        master_sheet.cell(row=m_row, column=5).value = "N/A"
                
                worker_wb.close()
                try:
                    os.remove(file_path)
                except Exception:
                    pass
                    
        master_wb.save(master_path)
